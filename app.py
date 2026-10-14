import os
import threading
import uuid
import io
import time
from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from prometheus_client import Counter, Histogram, generate_latest, REGISTRY
from models import db, User, Article, Chapter
from generator import get_genres
from ai_generator import generate_article_ai, generate_long_article_ai, test_connection, generate_outline_ai, generate_single_chapter_ai
from config import get_config, setup_logging

app = Flask(__name__)
app.config.from_object(get_config())


def safe_get(data, key, default='', coerce=None):
    """安全获取请求参数，防止 NoneType 和类型转换异常"""
    if data is None:
        return default
    val = data.get(key, default)
    if coerce is not None and val != default:
        try:
            val = coerce(val)
        except (ValueError, TypeError):
            val = default
    return val

# Prometheus 指标
REQUEST_COUNT = Counter(
    'flask_http_requests_total', 'Total HTTP requests',
    ['method', 'endpoint', 'status']
)
REQUEST_LATENCY = Histogram(
    'flask_http_request_duration_seconds', 'HTTP request latency',
    ['method', 'endpoint']
)

# 应用日志
setup_logging(app)

# 频率限制
limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["200 per day", "50 per hour"],
    storage_uri="memory://",
)


# Prometheus 请求中间件
@app.before_request
def before_request():
    request._start_time = time.time()


@app.after_request
def after_request(response):
    if hasattr(request, '_start_time'):
        elapsed = time.time() - request._start_time
        endpoint = request.endpoint or 'unknown'
        REQUEST_COUNT.labels(
            method=request.method,
            endpoint=endpoint,
            status=response.status_code
        ).inc()
        REQUEST_LATENCY.labels(
            method=request.method,
            endpoint=endpoint
        ).observe(elapsed)
    return response


db.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'index'


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


# ============================================================
# 页面路由
# ============================================================

@app.route('/')
def index():
    return render_template('index.html', genres=get_genres())


@app.route('/history')
def history_page():
    return render_template('history.html')


@app.route('/article/<int:article_id>')
def article_detail_page(article_id):
    return render_template('detail.html', article_id=article_id)


# ============================================================
# 用户认证 API
# ============================================================

@app.route('/api/register', methods=['POST'])
@limiter.limit("10 per minute")
def api_register():
    data = request.get_json(silent=True) or {}
    username = safe_get(data, 'username', '').strip()
    password = safe_get(data, 'password', '').strip()

    if not username or not password:
        return jsonify({'error': '用户名和密码不能为空'}), 400
    if len(username) < 2 or len(username) > 20:
        return jsonify({'error': '用户名长度需在2-20个字符之间'}), 400
    if len(password) < 6:
        return jsonify({'error': '密码长度至少6个字符'}), 400

    if User.query.filter_by(username=username).first():
        return jsonify({'error': '用户名已存在'}), 400

    user = User(username=username)
    user.set_password(password)
    db.session.add(user)
    db.session.commit()

    login_user(user)
    return jsonify({'message': '注册成功', 'username': user.username})


@app.route('/api/login', methods=['POST'])
@limiter.limit("20 per minute")
def api_login():
    data = request.get_json(silent=True) or {}
    username = safe_get(data, 'username', '').strip()
    password = safe_get(data, 'password', '').strip()

    user = User.query.filter_by(username=username).first()
    if not user or not user.check_password(password):
        return jsonify({'error': '用户名或密码错误'}), 401

    login_user(user)
    return jsonify({'message': '登录成功', 'username': user.username})


@app.route('/api/logout', methods=['POST'])
def api_logout():
    logout_user()
    return jsonify({'message': '已退出登录'})


@app.route('/api/user/status')
def api_user_status():
    if current_user.is_authenticated:
        return jsonify({'logged_in': True, 'username': current_user.username})
    return jsonify({'logged_in': False})


# ============================================================
# 文章生成 API
# ============================================================

# ============================================================
# AI 生成 API
# ============================================================

@app.route('/api/ai/test', methods=['GET'])
@login_required
def api_ai_test():
    result = test_connection()
    return jsonify({'result': result})


@app.route('/api/ai/generate', methods=['POST'])
@login_required
@limiter.limit("5 per minute")
def api_ai_generate():
    data = request.get_json(silent=True) or {}
    genre = safe_get(data, 'genre', '散文')
    topic = safe_get(data, 'topic', '').strip()
    keyword = safe_get(data, 'keyword', '').strip()
    length = safe_get(data, 'length', 'medium')
    style = safe_get(data, 'style', '抒情')
    perspective = safe_get(data, 'perspective', '第三人称')
    language_style = safe_get(data, 'language_style', '现代')
    constraints = safe_get(data, 'constraints', '').strip()
    format_type = safe_get(data, 'format_type', 'article')

    if not topic or not keyword:
        return jsonify({'error': '主题和关键词不能为空'}), 400

    try:
        content = generate_article_ai(genre, topic, keyword, length, style, perspective, language_style, constraints, format_type)
    except Exception as e:
        return jsonify({'error': f'AI 生成失败: {str(e)}'}), 500

    article = Article(
        title=f'{topic} - {genre}（AI）',
        genre=genre,
        topic=topic,
        keyword=keyword,
        length=length,
        word_count=len(content.replace('\n', '').replace(' ', '')),
        user_id=current_user.id if current_user.is_authenticated else None,
    )
    db.session.add(article)
    db.session.flush()

    chapter = Chapter(
        article_id=article.id,
        chapter_index=1,
        chapter_title='全文',
        content=content,
        word_count=len(content.replace('\n', '').replace(' ', '')),
    )
    db.session.add(chapter)
    db.session.commit()

    return jsonify({
        'article_id': article.id,
        'genre': genre,
        'topic': topic,
        'keyword': keyword,
        'length': length,
        'content': content,
    })


@app.route('/api/ai/generate/long', methods=['POST'])
@login_required
@limiter.limit("3 per minute")
def api_ai_generate_long():
    data = request.get_json(silent=True) or {}
    genre = safe_get(data, 'genre', '散文')
    topic = safe_get(data, 'topic', '').strip()
    keyword = safe_get(data, 'keyword', '').strip()
    target_words = safe_get(data, 'target_words', 10000, coerce=int)
    style = safe_get(data, 'style', '抒情')
    perspective = safe_get(data, 'perspective', '第三人称')
    language_style = safe_get(data, 'language_style', '现代')
    constraints = safe_get(data, 'constraints', '').strip()
    format_type = safe_get(data, 'format_type', 'article')

    if not topic or not keyword:
        return jsonify({'error': '主题和关键词不能为空'}), 400

    target_words = min(max(target_words, 3000), 500000)

    try:
        chapters_data = generate_long_article_ai(genre, topic, keyword, target_words, style, perspective, language_style, constraints, format_type)
    except Exception as e:
        return jsonify({'error': f'AI 生成失败: {str(e)}'}), 500

    total_words = sum(ch['word_count'] for ch in chapters_data)
    article = Article(
        title=f'{topic} - {genre}（AI长篇）',
        genre=genre,
        topic=topic,
        keyword=keyword,
        length='long',
        word_count=total_words,
        user_id=current_user.id if current_user.is_authenticated else None,
    )
    db.session.add(article)
    db.session.flush()

    for ch_data in chapters_data:
        chapter = Chapter(
            article_id=article.id,
            chapter_index=ch_data['chapter_index'],
            chapter_title=ch_data['chapter_title'],
            content=ch_data['content'],
            word_count=ch_data['word_count'],
        )
        db.session.add(chapter)

    db.session.commit()

    return jsonify({
        'article_id': article.id,
        'genre': genre,
        'topic': topic,
        'keyword': keyword,
        'target_words': target_words,
        'actual_words': total_words,
        'chapter_count': len(chapters_data),
        'chapters': chapters_data,
    })


# ============================================================
# 大纲与分步生成 API
# ============================================================

@app.route('/api/ai/outline', methods=['POST'])
@login_required
@limiter.limit("5 per minute")
def api_ai_outline():
    data = request.get_json(silent=True) or {}
    genre = safe_get(data, 'genre', '散文')
    topic = safe_get(data, 'topic', '').strip()
    keyword = safe_get(data, 'keyword', '').strip()
    target_words = safe_get(data, 'target_words', 10000, coerce=int)
    style = safe_get(data, 'style', '抒情')
    perspective = safe_get(data, 'perspective', '第三人称')
    language_style = safe_get(data, 'language_style', '现代')
    constraints = safe_get(data, 'constraints', '').strip()

    if not topic or not keyword:
        return jsonify({'error': '主题和关键词不能为空'}), 400

    target_words = min(max(target_words, 500), 2000000)

    try:
        from ai_generator import generate_outline_ai
        outline = generate_outline_ai(genre, topic, keyword, target_words, style, perspective, language_style, constraints)
    except Exception as e:
        return jsonify({'error': f'大纲生成失败: {str(e)}'}), 500

    return jsonify({
        'genre': genre,
        'topic': topic,
        'keyword': keyword,
        'target_words': target_words,
        'outline': outline,
    })


@app.route('/api/ai/generate/chapter', methods=['POST'])
@login_required
@limiter.limit("10 per minute")
def api_ai_generate_chapter():
    data = request.get_json(silent=True) or {}
    article_id = safe_get(data, 'article_id', None)
    chapter_index = safe_get(data, 'chapter_index', 1, coerce=int)
    chapter_title = safe_get(data, 'chapter_title', '')
    total_chapters = safe_get(data, 'total_chapters', 1, coerce=int)
    prev_ending = safe_get(data, 'prev_ending', '')
    genre = safe_get(data, 'genre', '散文')
    topic = safe_get(data, 'topic', '')
    keyword = safe_get(data, 'keyword', '')
    style = safe_get(data, 'style', '抒情')
    perspective = safe_get(data, 'perspective', '第三人称')
    language_style = safe_get(data, 'language_style', '现代')
    constraints = safe_get(data, 'constraints', '')

    try:
        from ai_generator import generate_single_chapter_ai
        content = generate_single_chapter_ai(
            genre, topic, keyword, chapter_index, total_chapters,
            chapter_title, prev_ending, style, perspective, language_style, constraints
        )
    except Exception as e:
        return jsonify({'error': f'章节生成失败: {str(e)}'}), 500

    word_count = len(content.replace('\n', '').replace(' ', ''))

    if article_id:
        article = Article.query.get_or_404(article_id)
        existing = Chapter.query.filter_by(article_id=article_id, chapter_index=chapter_index).first()
        if existing:
            existing.content = content
            existing.word_count = word_count
            existing.chapter_title = chapter_title
        else:
            ch = Chapter(
                article_id=article_id,
                chapter_index=chapter_index,
                chapter_title=chapter_title,
                content=content,
                word_count=word_count,
            )
            db.session.add(ch)
        total_words = sum(c.word_count for c in article.chapters)
        article.word_count = total_words
        db.session.commit()
        return jsonify({'article_id': article_id, 'chapter_index': chapter_index, 'chapter_title': chapter_title, 'content': content, 'word_count': word_count})
    else:
        return jsonify({'chapter_index': chapter_index, 'chapter_title': chapter_title, 'content': content, 'word_count': word_count})


@app.route('/api/articles/<int:article_id>/chapters/<int:chapter_id>/regenerate', methods=['POST'])
@login_required
def api_regenerate_chapter(article_id, chapter_id):
    data = request.get_json()
    article = Article.query.get_or_404(article_id)
    chapter = Chapter.query.get_or_404(chapter_id)

    prev_chapter = Chapter.query.filter_by(article_id=article_id, chapter_index=chapter.chapter_index - 1).first()
    prev_ending = prev_chapter.content[-100:] if prev_chapter and len(prev_chapter.content) > 100 else (prev_chapter.content if prev_chapter else '')

    total_chapters = len(article.chapters)

    try:
        from ai_generator import generate_single_chapter_ai
        content = generate_single_chapter_ai(
            article.genre, article.topic, article.keyword,
            chapter.chapter_index, total_chapters,
            chapter.chapter_title, prev_ending,
            data.get('style', '抒情'), data.get('perspective', '第三人称'),
            data.get('language_style', '现代'), data.get('constraints', '')
        )
    except Exception as e:
        return jsonify({'error': f'重新生成失败: {str(e)}'}), 500

    chapter.content = content
    chapter.word_count = len(content.replace('\n', '').replace(' ', ''))
    article.word_count = sum(c.word_count for c in article.chapters)
    db.session.commit()

    return jsonify({'chapter_id': chapter.id, 'content': content, 'word_count': chapter.word_count})


@app.route('/api/articles/<int:article_id>/chapters/<int:chapter_id>', methods=['PUT'])
@login_required
def api_edit_chapter(article_id, chapter_id):
    data = request.get_json()
    chapter = Chapter.query.get_or_404(chapter_id)
    article = Article.query.get_or_404(article_id)

    if 'content' in data:
        chapter.content = data['content']
        chapter.word_count = len(data['content'].replace('\n', '').replace(' ', ''))
    if 'chapter_title' in data:
        chapter.chapter_title = data['chapter_title']

    article.word_count = sum(c.word_count for c in article.chapters)
    db.session.commit()

    return jsonify({'message': '已更新', 'chapter_id': chapter.id, 'word_count': chapter.word_count})


@app.route('/api/articles/<int:article_id>/chapters', methods=['POST'])
@login_required
def api_add_chapter(article_id):
    data = request.get_json()
    article = Article.query.get_or_404(article_id)

    max_index = max([c.chapter_index for c in article.chapters], default=0)
    new_index = max_index + 1

    chapter_title = data.get('chapter_title', f'第{new_index}章')
    content = data.get('content', '')

    if not content:
        prev_chapter = Chapter.query.filter_by(article_id=article_id, chapter_index=max_index).first()
        prev_ending = prev_chapter.content[-100:] if prev_chapter and len(prev_chapter.content) > 100 else ''
        total_chapters = new_index
        try:
            from ai_generator import generate_single_chapter_ai
            content = generate_single_chapter_ai(
                article.genre, article.topic, article.keyword,
                new_index, total_chapters, chapter_title, prev_ending,
                data.get('style', '抒情'), data.get('perspective', '第三人称'),
                data.get('language_style', '现代'), data.get('constraints', '')
            )
        except Exception as e:
            return jsonify({'error': f'续写失败: {str(e)}'}), 500

    word_count = len(content.replace('\n', '').replace(' ', ''))
    chapter = Chapter(
        article_id=article_id,
        chapter_index=new_index,
        chapter_title=chapter_title,
        content=content,
        word_count=word_count,
    )
    db.session.add(chapter)
    article.word_count = sum(c.word_count for c in article.chapters) + word_count
    db.session.commit()

    return jsonify({'chapter_id': chapter.id, 'chapter_index': new_index, 'chapter_title': chapter_title, 'content': content, 'word_count': word_count})


# ============================================================
# 导出与评分 API
# ============================================================

@app.route('/api/articles/<int:article_id>/export/word', methods=['GET'])
@login_required
def api_export_word(article_id):
    article = Article.query.get_or_404(article_id)
    from docx import Document

    doc = Document()
    doc.add_heading(article.title, 0)

    for ch in article.chapters:
        doc.add_heading(ch.chapter_title, level=1)
        for para in ch.content.split('\n'):
            para = para.strip()
            if para:
                doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        as_attachment=True,
        download_name=f'{article.topic}_article_{article_id}.docx',
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )


@app.route('/api/articles/<int:article_id>/export/pdf', methods=['GET'])
@login_required
def api_export_pdf(article_id):
    article = Article.query.get_or_404(article_id)
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.units import mm

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    font_path = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
    try:
        pdfmetrics.registerFont(TTFont('WQY', font_path))
        font_name = 'WQY'
    except Exception:
        font_name = 'Helvetica'

    y = height - 30 * mm
    c.setFont(font_name, 18)
    c.drawString(30 * mm, y, article.title[:40])
    y -= 15 * mm

    c.setFont(font_name, 11)
    for ch in article.chapters:
        c.setFont(font_name, 14)
        y -= 10 * mm
        if y < 30 * mm:
            c.showPage()
            y = height - 30 * mm
        c.drawString(30 * mm, y, ch.chapter_title[:50])
        y -= 8 * mm

        c.setFont(font_name, 11)
        for line in ch.content.split('\n'):
            line = line.strip()
            if not line:
                y -= 5 * mm
                continue
            while len(line) > 38:
                if y < 25 * mm:
                    c.showPage()
                    y = height - 30 * mm
                c.drawString(30 * mm, y, line[:38])
                line = line[38:]
                y -= 6 * mm
            if y < 25 * mm:
                c.showPage()
                y = height - 30 * mm
            c.drawString(30 * mm, y, line)
            y -= 6 * mm

    c.save()
    buf.seek(0)

    return send_file(
        buf,
        as_attachment=True,
        download_name=f'{article.topic}_article_{article_id}.pdf',
        mimetype='application/pdf'
    )


@app.route('/api/articles/<int:article_id>/score', methods=['GET'])
@login_required
def api_article_score(article_id):
    article = Article.query.get_or_404(article_id)

    total_words = article.word_count
    chapter_count = len(article.chapters)
    avg_chapter_words = total_words // chapter_count if chapter_count > 0 else 0

    avg_paragraph_len = 0
    total_paragraphs = 0
    for ch in article.chapters:
        paras = [p.strip() for p in ch.content.split('\n') if p.strip()]
        total_paragraphs += len(paras)
        if paras:
            avg_paragraph_len += sum(len(p) for p in paras)
    avg_paragraph_len = avg_paragraph_len // total_paragraphs if total_paragraphs > 0 else 0

    has_dialogue = any('：' in ch.content or '"' in ch.content or '\u201c' in ch.content for ch in article.chapters)
    has_bold = any('**' in ch.content for ch in article.chapters)

    scores = {
        'word_count': min(100, total_words // 30) if total_words > 0 else 0,
        'structure': min(100, chapter_count * 20),
        'paragraph': min(100, avg_paragraph_len * 2) if avg_paragraph_len > 0 else 0,
        'dialogue': 80 if has_dialogue else 30,
        'formatting': 70 if has_bold else 40,
    }
    overall = sum(scores.values()) // len(scores)

    return jsonify({
        'overall': overall,
        'scores': scores,
        'stats': {
            'total_words': total_words,
            'chapter_count': chapter_count,
            'avg_chapter_words': avg_chapter_words,
            'total_paragraphs': total_paragraphs,
            'avg_paragraph_len': avg_paragraph_len,
            'has_dialogue': has_dialogue,
            'has_bold': has_bold,
        },
        'suggestions': [
            '增加字数以丰富内容' if total_words < 1000 else '字数充足',
            '增加章节结构' if chapter_count < 3 else '结构完整',
            '增加对话描写' if not has_dialogue else '对话丰富',
            '增加金句加粗' if not has_bold else '格式规范',
            '段落过长，建议拆分' if avg_paragraph_len > 300 else '段落适中',
        ]
    })


# ============================================================
# 文章管理 API
# ============================================================

@app.route('/api/articles', methods=['GET'])
@login_required
def api_article_list():
    query = Article.query
    if current_user.is_authenticated:
        query = query.filter_by(user_id=current_user.id)
    else:
        query = query.filter_by(user_id=None)

    keyword = (request.args.get('q') or '').strip()
    genre = (request.args.get('genre') or '').strip()
    fav = (request.args.get('fav') or '').strip()

    if keyword:
        like = f'%{keyword}%'
        query = query.filter(db.or_(Article.title.like(like), Article.topic.like(like), Article.keyword.like(like)))
    if genre:
        query = query.filter_by(genre=genre)
    if fav == '1':
        query = query.filter_by(is_favorite=True)

    articles = query.order_by(Article.created_at.desc()).all()
    return jsonify({'articles': [a.to_dict() for a in articles]})


@app.route('/api/articles/<int:article_id>/rename', methods=['PUT'])
@login_required
def api_article_rename(article_id):
    data = request.get_json(silent=True) or {}
    title = (safe_get(data, 'title', '') or '').strip()
    if not title:
        return jsonify({'error': '标题不能为空'}), 400
    article = Article.query.get_or_404(article_id)
    article.title = title[:200]
    db.session.commit()
    return jsonify({'message': '已重命名', 'title': article.title})


@app.route('/api/articles/<int:article_id>/favorite', methods=['POST'])
@login_required
def api_article_favorite(article_id):
    article = Article.query.get_or_404(article_id)
    article.is_favorite = not bool(article.is_favorite)
    db.session.commit()
    return jsonify({'message': '已更新', 'is_favorite': bool(article.is_favorite)})


@app.route('/api/articles/<int:article_id>', methods=['GET'])
@login_required
def api_article_detail(article_id):
    article = Article.query.get_or_404(article_id)
    return jsonify({
        'article': article.to_dict(),
        'chapters': [ch.to_dict() for ch in article.chapters],
    })


@app.route('/api/articles/<int:article_id>', methods=['DELETE'])
@login_required
def api_article_delete(article_id):
    article = Article.query.get_or_404(article_id)
    db.session.delete(article)
    db.session.commit()
    return jsonify({'message': '已删除'})


@app.route('/api/genres', methods=['GET'])
def api_genres():
    return jsonify({'genres': get_genres()})


@app.route('/health')
def health():
    import time
    db_ok = True
    db_error = None
    try:
        db.session.execute(db.text('SELECT 1'))
    except Exception as e:
        db_ok = False
        db_error = str(e)
    return jsonify({
        'status': 'healthy' if db_ok else 'degraded',
        'database': 'ok' if db_ok else 'error',
        'database_error': db_error,
        'timestamp': int(time.time()),
    }), 200 if db_ok else 503


@app.route('/metrics')
def metrics():
    """Prometheus 指标端点"""
    return generate_latest(REGISTRY), 200, {'Content-Type': 'text/plain; charset=utf-8'}


# ============================================================
# 异步任务管理
# ============================================================

generation_tasks = {}


def _generate_from_outline(genre, topic, keyword, outline, style, perspective, language_style, constraints, format_type):
    """按照用户确认过的大纲逐节生成正文"""
    from ai_generator import generate_single_chapter_ai
    chapters = []
    prev_ending = ''
    total = len(outline)
    for idx, item in enumerate(outline, start=1):
        title = item.get('chapter_title', f'{topic}第{idx}节')
        print(f'[ASYNC] Generating {idx}/{total}: {title}', flush=True)
        content = generate_single_chapter_ai(
            genre, topic, keyword, idx, total, title, prev_ending,
            style, perspective, language_style, constraints, format_type
        )
        content = content.strip()
        word_count = len(content.replace('\n', '').replace(' ', ''))
        prev_ending = content[-100:] if len(content) > 100 else content
        chapters.append({
            'chapter_index': idx,
            'chapter_title': title,
            'content': content,
            'word_count': word_count,
        })
        print(f'[ASYNC] {idx}/{total} done: {word_count} words', flush=True)
    return chapters


def _run_long_generation(task_id, genre, topic, keyword, target_words, style, perspective, language_style, constraints, is_ai, format_type='article', custom_outline=None):
    """后台线程执行长文生成（支持传入用户确认过的大纲）"""
    import sys
    task = generation_tasks[task_id]
    try:
        task['status'] = 'generating'
        print(f'[ASYNC] Task {task_id} started: genre={genre}, topic={topic}, target_words={target_words}, is_ai={is_ai}, custom_outline={bool(custom_outline)}', flush=True)

        if is_ai:
            if custom_outline:
                chapters_data = _generate_from_outline(
                    genre, topic, keyword, custom_outline, style, perspective, language_style, constraints, format_type
                )
            else:
                chapters_data = generate_long_article_ai(genre, topic, keyword, target_words, style, perspective, language_style, constraints, format_type)
        else:
            chapters_data = generate_long_article(genre, topic, keyword, target_words)

        print(f'[ASYNC] Task {task_id}: generation done, chapters={len(chapters_data)}', flush=True)

        total_chapters = len(chapters_data)
        total_words = 0
        saved_chapters = []

        for i, ch_data in enumerate(chapters_data):
            if task.get('cancelled'):
                task['status'] = 'cancelled'
                return
            task['current_chapter'] = i + 1
            task['total_chapters'] = total_chapters
            task['progress'] = int((i + 1) / total_chapters * 100)
            total_words += ch_data['word_count']
            saved_chapters.append(ch_data)
            print(f'[ASYNC] Task {task_id}: chapter {i+1}/{total_chapters} done', flush=True)

        with app.app_context():
            article = Article(
                title=f'{topic} - {genre}（{"AI长篇" if is_ai else "长篇"}）',
                genre=genre,
                topic=topic,
                keyword=keyword,
                length='long',
                word_count=total_words,
                user_id=task.get('user_id'),
            )
            db.session.add(article)
            db.session.flush()

            for ch_data in saved_chapters:
                chapter = Chapter(
                    article_id=article.id,
                    chapter_index=ch_data['chapter_index'],
                    chapter_title=ch_data['chapter_title'],
                    content=ch_data['content'],
                    word_count=ch_data['word_count'],
                )
                db.session.add(chapter)

            db.session.commit()

            task['status'] = 'completed'
            task['article_id'] = article.id
            task['actual_words'] = total_words
            task['chapter_count'] = total_chapters
            print(f'[ASYNC] Task {task_id}: completed, article_id={article.id}', flush=True)

    except Exception as e:
        task['status'] = 'failed'
        task['error'] = str(e)
        print(f'[ASYNC] Task {task_id}: FAILED: {e}', flush=True)
        import traceback
        traceback.print_exc(file=sys.stdout)


@app.route('/api/generate/async', methods=['POST'])
@login_required
@limiter.limit("3 per minute")
def api_generate_async():
    data = request.get_json()
    genre = data.get('genre', '散文')
    topic = data.get('topic', '').strip()
    keyword = data.get('keyword', '').strip()
    target_words = int(data.get('target_words', 10000))
    is_ai = data.get('is_ai', True)
    style = data.get('style', '抒情')
    perspective = data.get('perspective', '第三人称')
    language_style = data.get('language_style', '现代')
    constraints = data.get('constraints', '').strip()
    format_type = data.get('format_type', 'article')
    custom_outline = data.get('custom_outline')

    if not topic or not keyword:
        return jsonify({'error': '主题和关键词不能为空'}), 400

    target_words = min(max(target_words, 500), 2000000)

    task_id = str(uuid.uuid4())
    generation_tasks[task_id] = {
        'status': 'pending',
        'progress': 0,
        'current_chapter': 0,
        'total_chapters': 0,
        'article_id': None,
        'actual_words': 0,
        'chapter_count': 0,
        'error': None,
        'cancelled': False,
        'user_id': current_user.id if current_user.is_authenticated else None,
    }

    thread = threading.Thread(
        target=_run_long_generation,
        args=(task_id, genre, topic, keyword, target_words, style, perspective, language_style, constraints, is_ai, format_type, custom_outline)
    )
    thread.daemon = True
    thread.start()

    return jsonify({'task_id': task_id, 'status': 'pending'})


@app.route('/api/generate/status/<task_id>', methods=['GET'])
@login_required
def api_generate_status(task_id):
    task = generation_tasks.get(task_id)
    if not task:
        return jsonify({'error': '任务不存在'}), 404

    result = {
        'status': task['status'],
        'progress': task['progress'],
        'current_chapter': task['current_chapter'],
        'total_chapters': task['total_chapters'],
    }

    if task['status'] == 'completed':
        result['article_id'] = task['article_id']
        result['actual_words'] = task['actual_words']
        result['chapter_count'] = task['chapter_count']
    elif task['status'] == 'failed':
        result['error'] = task['error']

    return jsonify(result)


@app.route('/api/generate/cancel/<task_id>', methods=['POST'])
@login_required
def api_generate_cancel(task_id):
    task = generation_tasks.get(task_id)
    if not task:
        return jsonify({'error': '任务不存在'}), 404
    task['cancelled'] = True
    return jsonify({'message': '已请求取消'})


# ============================================================
# 全局错误处理
# ============================================================

@app.errorhandler(404)
def not_found_error(error):
    return render_template('404.html'), 404


@app.errorhandler(500)
def internal_error(error):
    db.session.rollback()
    return render_template('500.html'), 500


with app.app_context():
    db.create_all()


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=os.environ.get('FLASK_ENV') == 'development')
